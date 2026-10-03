"""End-to-end test celej aplikácie bez siete: falošný Telegram (zbiera správy), lokálna
tabuľka (xlsx na disku), falošný CalDAV zoznam, riadený čas. Simuluje reálne dni."""
import asyncio
import io
from dataclasses import replace
from datetime import date, datetime, time as dtime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from trener import app as A
from trener.config import Config
from trener.model import EVENING, MORNING, Day
from trener.reminders import ReminderSync
from trener.smb_io import LocalBackend
from trener.store import Store
from trener.table import COL, SHEET_DAYS, SHEET_SETTINGS
from test_reminders import FakeTodos


class FakeCalendar(FakeTodos):
    """Falošný iCloud kalendár – rovnaké rozhranie ako EventCalendar, len v pamäti."""
    component = "VEVENT"

    def events(self):
        return sorted(self.list(), key=lambda i: i.summary)

TZ = ZoneInfo("Europe/Bratislava")
# Pozor pri písaní testov: D je STREDA – teda deň s rannou aj večernou fázou.
# Utorok/štvrtok/sobota majú len ráno a nedeľa je voľno; na tie treba iné dátumy.
D = date(2026, 9, 2)          # streda
D_UT = date(2026, 9, 1)       # utorok – len ráno
D_NE = date(2026, 9, 6)       # nedeľa – voľno
D_PO = date(2026, 9, 7)       # pondelok – nové X
# seed_x=3 → streda 2*3 ráno + 2*3 večer = 12, rovnako ako starý seed_goal=12


def at(h, m=0, s=0, d=D):
    return datetime.combine(d, dtime(h, m, s), tzinfo=TZ)


class Harness:
    def __init__(self, tmp: Path, start: datetime, mode: str = "caldav"):
        self.sent: list[str] = []
        self.clock = [start]
        self.mode = mode
        cfg = Config(bot_token="x", owner_chat_id=1, tz=TZ, state_db=tmp / "t.db", log_file=None,
                     table_backend="local", smb_server="", smb_share="", smb_path="", smb_username="",
                     smb_password="", local_table_path=tmp / "kliky.xlsx", table_sync_seconds=120,
                     calendar_mode="off", calendar_url="https://caldav.icloud.com/", calendar_user=None,
                     calendar_password=None, calendar_name="Kliky", calendar_color="#FF6B35",
                     calendar_minutes=25, calendar_alarms=(0, 3, 7, 12, 20), calendar_sync_seconds=120,
                     reminders_mode="caldav", reminders_list="Kliky", shortcut_token="tok",
                     caldav_url=None, caldav_username=None, caldav_password=None, caldav_list="Kliky",
                     reminders_sync_seconds=120, web_port=0, wake_token=None,
                     alarm_mode="off", pushover_token=None, pushover_user=None, pushover_device=None,
                     alarm_priority=2, alarm_sound="persistent", alarm_retry=60, alarm_expire=600,
                     alarm_webhook_url=None,
                     seed_x=3, seed_x_step=1, seed_morning="07:00", seed_evening="19:20", tick_seconds=30)
        self.cfg = cfg
        self.store = Store(cfg.state_db)
        self.backend = LocalBackend(cfg.local_table_path)
        self.todos = FakeTodos()

        if mode == "shortcuts":
            cfg = replace(cfg, reminders_mode="shortcuts")
        self.cfg = cfg
        self.cal = FakeCalendar()
        if mode == "calendar":
            cfg = replace(cfg, calendar_mode="icloud", reminders_mode="off")
            self.cfg = cfg
        self.t = self._make_trainer()
        self.t.bootstrap()

    def _make_trainer(self):
        async def send(text):
            self.sent.append(text)

        t = A.Trainer(self.cfg, self.store, self.backend, None, send,
                      self.cal if self.mode == "calendar" else None)
        if self.mode == "shortcuts":
            t.rem = None
            t.todos = None
        else:
            t.make_reminders = False
            t.todos = self.todos
            t.rem = ReminderSync(self.todos, self.store, TZ)
        t.now = lambda: self.clock[0]
        return t

    def restart(self, when: datetime):
        """Reštart procesu: nový Trainer aj nové spojenie do tej istej DB a tabuľky."""
        self.clock[0] = when
        self.store.close()
        self.store = Store(self.cfg.state_db)
        self.t = self._make_trainer()
        self.t.bootstrap()
        return self.t

    def tick_at(self, when: datetime):
        self.clock[0] = when
        asyncio.run(self.t.tick())

    def text(self, when: datetime, msg: str) -> str:
        self.clock[0] = when
        return asyncio.run(self.t.handle_text(msg))

    def take(self) -> list[str]:
        out, self.sent = self.sent, []
        return out

    def table_rows(self):
        wb = load_workbook(self.cfg.local_table_path, data_only=True)
        ws = wb[SHEET_DAYS]
        return {ws.cell(r, 1).value.date() if hasattr(ws.cell(r, 1).value, "date") else ws.cell(r, 1).value:
                (ws.cell(r, COL["Cieľ"]).value, ws.cell(r, COL["Ráno"]).value, ws.cell(r, COL["Večer"]).value,
                 ws.cell(r, COL["Stav"]).value) for r in range(2, ws.max_row + 1)}

    def edit_table(self, d: date, **cols):
        wb = load_workbook(self.cfg.local_table_path)
        ws = wb[SHEET_DAYS]
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if (v.date() if hasattr(v, "date") else v) == d:
                for k, val in cols.items():
                    ws.cell(r, COL[k], val)
        wb.save(self.cfg.local_table_path)

    def edit_setting(self, label: str, value):
        wb = load_workbook(self.cfg.local_table_path)
        ws = wb[SHEET_SETTINGS]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value == label:
                ws.cell(r, 2, value)
        wb.save(self.cfg.local_table_path)

    def reminder(self, word):
        items = self.todos.by_summary_contains(word)
        return items[0] if items else None


@pytest.fixture
def h(tmp_path):
    return Harness(tmp_path, at(6, 0))


@pytest.fixture
def hc(tmp_path):
    """Harness s budíkom v iCloud kalendári."""
    return Harness(tmp_path, at(6, 0), mode="calendar")


@pytest.fixture
def hs(tmp_path):
    """Harness v režime iCloud Pripomienok cez iOS Skratku."""
    return Harness(tmp_path, at(6, 0), mode="shortcuts")


def test_first_start_creates_table_and_reminders(h):
    h.tick_at(at(6, 0))
    assert h.cfg.local_table_path.exists()
    rows = h.table_rows()
    assert rows[D][0] == 12 and rows[D][3].startswith("⏳")
    assert h.reminder("Ráno").summary == "💪 Ráno: 6 klikov"
    assert h.reminder("Večer").summary == "💪 Večer: 6 klikov"
    assert h.take() == []          # žiadna správa len tak


def test_user_scenario_two_plus_two(h):
    h.tick_at(at(6, 0))
    r1 = h.text(at(7, 5), "2")
    assert "Ráno: 2/6" in r1 and "Dnes: 2/12" in r1 and "🎉" not in r1
    r2 = h.text(at(7, 20), "ráno som dal ďalšie 2")
    assert "Ráno: 4/6" in r2 and "Dnes: 4/12" in r2 and "zostáva 8" in r2 and "🎉" not in r2
    h.tick_at(at(7, 21))
    assert h.table_rows()[D][1:3] == (4, 0)
    assert h.reminder("Ráno").summary == "💪 Ráno: 2 klikov"
    assert h.reminder("Večer").summary == "💪 Večer: 8 klikov"


def test_nag_limits_and_stop_on_done(h):
    h.tick_at(at(6, 0))
    for m in range(0, 121, 5):
        h.tick_at(at(7, 0) + timedelta(minutes=m))
    nags = [s for s in h.take() if s.startswith("☀️")]
    assert len(nags) == 3 and "(1/3)" in nags[0] and "(3/3)" in nags[2]
    # večer: prvá výzva o 19:20, po hlásení celého zvyšku ticho
    h.tick_at(at(19, 20))
    assert [s[0] for s in h.take()] == ["🌙"]
    reply = h.text(at(19, 30), "12")
    assert "🎉" in reply
    for m in range(35, 120, 5):
        h.tick_at(at(19, m) if m < 60 else at(20, m - 60))
    assert h.take() == []
    assert h.reminder("Ráno").completed and h.reminder("Večer").completed


def test_freeze_blocks_nags_and_new_reminders(h):
    h.tick_at(at(6, 0))
    msg = asyncio.run(h.t.set_frozen(True))
    assert "Zamrazené" in msg
    h.tick_at(at(7, 0)); h.tick_at(at(7, 30)); h.tick_at(at(19, 20))
    assert h.take() == []
    wb = load_workbook(h.cfg.local_table_path, data_only=True)
    ws = wb[SHEET_SETTINGS]
    vals = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    assert vals["Zamrazené"] == "ÁNO"
    # nový deň počas zamrazenia: riadok vznikne, pripomienky nie
    h.tick_at(at(0, 0, 30, D + timedelta(days=1)))
    assert D + timedelta(days=1) in h.table_rows()
    assert all(i.due.astimezone(TZ).date() == D for i in h.todos.list())
    assert h.take() == []            # ani polnočný verdikt pri zamrazení
    asyncio.run(h.t.set_frozen(False))
    h.tick_at(at(0, 2, 0, D + timedelta(days=1)))
    assert any(i.due.astimezone(TZ).date() == D + timedelta(days=1) for i in h.todos.list())


def test_freeze_from_table_cell(h):
    h.tick_at(at(6, 0))
    h.edit_setting("Zamrazené", "ÁNO")
    h.tick_at(at(6, 3))
    assert h.t.settings().frozen is True and h.store.get_day(D).frozen is True
    h.tick_at(at(7, 0))
    assert h.take() == []
    # odmrazenie cez tabuľku odmrazí aj dnešok → výzvy a pripomienky idú
    h.edit_setting("Zamrazené", "NIE")
    h.tick_at(at(7, 5))
    assert h.t.settings().frozen is False and h.store.get_day(D).frozen is False
    assert h.table_rows()[D][3].startswith("⏳")
    h.tick_at(at(7, 6))
    assert [m.startswith("☀️") for m in h.take()] == [True]


def test_yesterday_report_and_wake_floor(h):
    h.tick_at(at(6, 0))
    nd = D + timedelta(days=1)
    h.tick_at(at(0, 0, 30, nd))
    h.take()
    reply = h.text(at(8, 0, 0, nd), "včera večer 12")
    assert "Včera: 12/12" in reply and "dodatočne" in reply
    assert h.store.get_day(D).done and h.store.get_day(nd).total == 0
    # wake pred 04:00 sa ignoruje, po 04:00 spustí ranné výzvy hneď
    h.clock[0] = at(3, 30, 0, nd); h.t.wake(); assert h.t.woke_date is None
    h.clock[0] = at(5, 0, 0, nd); h.t.wake(); assert h.t.woke_date == nd
    h.tick_at(at(5, 0, 5, nd))
    assert [m.startswith("☀️") for m in h.take()] == [True]


def test_user_edits_table_counts_and_completes(h):
    h.tick_at(at(6, 0))
    h.edit_table(D, Ráno=6)
    h.tick_at(at(6, 3))
    day = h.store.get_day(D)
    assert day.morning == 6
    msgs = h.take()
    assert len(msgs) == 1 and msgs[0].startswith("📋") and "ráno 0 → 6" in msgs[0]
    assert h.reminder("Ráno").completed and not h.reminder("Večer").completed
    h.edit_table(D, Večer=6)
    h.tick_at(at(6, 6))
    msgs = h.take()
    assert len(msgs) == 1 and msgs[0].startswith("🎉") and "12/12" in msgs[0]
    assert h.reminder("Večer").completed
    h.tick_at(at(6, 9)); h.tick_at(at(7, 0)); h.tick_at(at(19, 20))
    assert h.take() == []            # splnenie sa oznámi len raz, výzvy žiadne
    assert h.table_rows()[D][3].startswith("✅")


def test_rollover_progression_and_failure_summary(h):
    h.tick_at(at(6, 0))
    h.text(at(8, 0), "12")
    h.take()
    nd = D + timedelta(days=1)       # štvrtok – len ráno 2X, X sa uprostred týždňa nemení
    h.tick_at(at(0, 0, 30, nd))
    assert h.store.get_day(nd).goal == 6
    assert h.take() == []            # splnený deň → žiadny verdikt
    zajtrajsie = [i.summary for i in h.todos.list() if i.due.astimezone(TZ).date() == nd]
    assert zajtrajsie == ["💪 Ráno: 6 klikov"]    # vo štvrtok večer nie je v pláne nič
    # ďalší deň bez klikov → jedna správa o nesplnení, týždeň je tým pokazený
    nd2 = nd + timedelta(days=1)     # piatok – ráno aj večer
    h.tick_at(at(0, 0, 30, nd2))
    msgs = h.take()
    assert len(msgs) == 1 and msgs[0].startswith("❌") and "0/6" in msgs[0]
    assert "Týždeň je tým pokazený" in msgs[0] and "Dnes (piatok): 6 ráno + 6 večer" in msgs[0]
    assert h.store.get_day(nd2).goal == 12
    # staré otvorené pripomienky preč, splnené z 2.9. ostali
    dues = sorted({i.due.astimezone(TZ).date() for i in h.todos.list()})
    assert dues == [D, nd2]


def test_downtime_days_become_frozen_and_no_nag_burst(h):
    h.tick_at(at(6, 0))
    later = at(11, 0, 0, D + timedelta(days=3))
    h.tick_at(later)
    days = {d.date: d for d in h.store.all_days()}
    assert days[D + timedelta(days=1)].frozen and days[D + timedelta(days=2)].frozen
    assert not days[D + timedelta(days=3)].frozen
    assert [m for m in h.take() if m.startswith("☀️")] == []
    assert h.store.get_day(D).frozen is False      # dnešok (2.9.) sa už hodnotí ako nesplnený
    assert any(m.startswith("❌") for m in h.sent) is False


def test_user_ticks_and_renames_in_reminders(h):
    h.tick_at(at(6, 0))
    m = h.reminder("Ráno")
    h.todos.user_edit(m.href, complete=True)
    h.tick_at(at(6, 3))
    assert h.store.get_day(D).morning == 6
    msgs = h.take()
    assert len(msgs) == 1 and msgs[0].startswith("📱")
    e = h.reminder("Večer")
    h.todos.user_edit(e.href, summary="Kliky večer: 6 kusov")
    h.tick_at(at(6, 6))
    assert h.t.settings().evening_title == "Kliky večer: {n} kusov"
    msgs = h.take()
    assert len(msgs) == 1 and "nový názov" in msgs[0]
    h.todos.user_edit(h.reminder("Kliky večer").href, due=at(20, 0))
    h.tick_at(at(6, 9))
    assert h.t.settings().evening_time == "20:00"
    msgs = h.take()
    assert len(msgs) == 1 and "nový čas 20:00" in msgs[0]
    wb = load_workbook(h.cfg.local_table_path, data_only=True)
    ws = wb[SHEET_SETTINGS]
    vals = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    assert vals["Večerný čas"] == "20:00" and vals["Názov večernej pripomienky"] == "Kliky večer: {n} kusov"
    # večerná výzva teraz o 20:00, nie 19:20
    h.tick_at(at(19, 20)); assert h.take() == []
    h.tick_at(at(20, 0)); assert [s[0] for s in h.take()] == ["🌙"]


def test_commands(h):
    h.tick_at(at(6, 0))
    msg = asyncio.run(h.t.set_x(10))            # /ciel 10 = X, nie cieľ dňa
    assert "X = 10 od tohto týždňa" in msg and "Dnes (streda): 20 ráno + 20 večer" in msg
    assert h.store.get_day(D).goal == 40 and h.t.settings().x == 10
    h.tick_at(at(6, 1))
    assert h.reminder("Ráno").summary == "💪 Ráno: 20 klikov"
    assert "06:30" in asyncio.run(h.t.set_time(MORNING, "06:30"))
    h.tick_at(at(6, 2))
    assert h.reminder("Ráno").due.astimezone(TZ).strftime("%H:%M") == "06:30"
    fix = asyncio.run(h.t.fix(EVENING, 3))
    assert "Večer: 3/20" in fix
    st = asyncio.run(h.t.status_text())
    assert "3/40" in st and "X = 10 → 20 ráno + 20 večer" in st and "Streak" in st
    assert "2" in asyncio.run(h.t.set_x_step(2)) and h.t.settings().x_step == 2
    assert "Tabuľka" in asyncio.run(h.t.table_info())
    assert "Sync hotový" in asyncio.run(h.t.force_sync())


def test_table_corrupt_is_not_overwritten(h):
    h.tick_at(at(6, 0))
    h.cfg.local_table_path.write_bytes(b"rozpisany subor")
    h.tick_at(at(6, 3))
    assert h.cfg.local_table_path.read_bytes() == b"rozpisany subor"
    assert h.t.table_ok is False
    for i in range(5):
        h.tick_at(at(6, 5 + 2 * i))
    assert any("nedostupná" in m for m in h.take())


def test_reminder_backend_down_does_not_block_nags(h):
    h.tick_at(at(6, 0))
    original = h.todos.list

    def boom():
        raise RuntimeError("radicale down")
    h.todos.list = boom
    h.tick_at(at(7, 0))
    assert [m.startswith("☀️") for m in h.take()] == [True]
    assert h.t.rem_ok is False
    h.todos.list = original
    h.tick_at(at(7, 2))
    assert h.t.rem_ok is True


def test_excel_holding_file_is_not_an_outage(h):
    from trener.smb_io import Busy
    h.tick_at(at(6, 0))
    real_write = h.backend.write

    def busy(data, expected):
        raise Busy("excel")
    h.backend.write = busy
    h.text(at(7, 5), "2")
    for m in range(6, 40, 1):
        h.tick_at(at(7, m))
    assert not any("nedostupná" in x for x in h.take())
    assert h.t.table_ok and h.t.table_busy and h.t.table_dirty
    assert "Excel" in asyncio.run(h.t.table_info())
    h.backend.write = real_write
    h.tick_at(at(7, 45))
    assert not h.t.table_busy and h.table_rows()[D][1] == 2


def test_no_table_write_churn(h):
    h.tick_at(at(6, 0))
    m1 = h.cfg.local_table_path.stat().st_mtime_ns
    for i in range(1, 8):
        h.tick_at(at(6, 2 * i, 5))
    assert h.cfg.local_table_path.stat().st_mtime_ns == m1


def test_stale_message_from_yesterday_ignored(h):
    h.tick_at(at(6, 0))
    reply = asyncio.run(h.t.handle_text("5", when=at(21, 0, 0, D - timedelta(days=1))))
    assert "nepočítam" in reply and h.store.get_day(D).total == 0
    reply = asyncio.run(h.t.handle_text("5", when=at(7, 0)))
    assert "Dnes: 5/12" in reply


def test_delayed_same_day_message_uses_send_time_for_session(h):
    h.tick_at(at(6, 0))
    # odoslané 18:45 (ranná fáza), spracované 19:30 po reštarte → ráno, nie večer
    reply = asyncio.run(h.t.handle_text("2", when=at(18, 45)))
    assert "Ráno: 2/6" in reply
    h.clock[0] = at(19, 30)
    reply = asyncio.run(h.t.handle_text("3", when=at(18, 50)))
    assert "Ráno: 5/6" in reply and h.store.get_day(D).evening == 0


def test_message_just_before_midnight_counts_for_yesterday(h):
    h.tick_at(at(6, 0))
    nd = D + timedelta(days=1)
    h.tick_at(at(0, 0, 30, nd)); h.take()
    reply = asyncio.run(h.t.handle_text("12", when=at(23, 59, 50)))   # spracované 00:00:30
    assert "Včera: 12/12" in reply and h.store.get_day(D).done and h.store.get_day(nd).total == 0


def test_timeout_keeps_inflight_until_thread_finishes(h):
    import threading
    h.tick_at(at(6, 0))
    gate = threading.Event()
    real_list = h.todos.list

    def slow_list():
        gate.wait(5)
        return real_list()
    h.todos.list = slow_list
    orig = h.t._guarded

    async def fast_guarded(flag, fn, *args, timeout):
        return await orig(flag, fn, *args, timeout=0.2)
    h.t._guarded = fast_guarded

    async def scenario():
        h.t.rem_dirty = True
        await h.t._sync_reminders()                 # timeout po 0.2 s, vlákno beží ďalej
        assert h.t._rem_inflight is True and h.t.rem_ok is False
        await h.t._sync_reminders()                 # druhá kópia sa nespustí
        assert h.t._rem_inflight is True
        gate.set()
        await asyncio.sleep(0.5)                    # vlákno dobehlo → zámok sa uvoľnil
        assert h.t._rem_inflight is False
        h.t._guarded = orig
        h.todos.list = real_list
        await h.t._sync_reminders()
        assert h.t.rem_ok is True and h.t._rem_inflight is False
    asyncio.run(scenario())


def test_undo_and_it_was_a_goal(h):
    h.tick_at(at(6, 0))
    reply = h.text(at(7, 5), "10")            # používateľ chcel X, bot zapísal kliky
    assert "Ráno: 10/6" in reply
    h.tick_at(at(7, 6))
    assert h.reminder("Ráno").completed        # ranná fáza „splnená“
    assert h.t.undo_available() == (True, True)
    msg = asyncio.run(h.t.undo_last(as_goal=True))
    assert "X = 10 od tohto týždňa" in msg and "Dnes (streda): 20 ráno + 20 večer" in msg
    day = h.store.get_day(D)
    assert (day.goal, day.morning, day.evening) == (40, 0, 0)      # 10 bolo X, teda 4X na stredu
    h.tick_at(at(7, 7))
    m = h.reminder("Ráno")
    assert not m.completed and m.summary == "💪 Ráno: 20 klikov"    # odčiarknutie zrušené, nový cieľ
    assert h.table_rows()[D][0] == 40 and h.table_rows()[D][1] == 0
    assert h.t.undo_available() == (False, False)
    # obyčajné vrátenie
    h.text(at(7, 10), "3")
    assert asyncio.run(h.t.undo_last()).startswith("↩️")
    assert h.store.get_day(D).morning == 0
    assert "Nie je čo vrátiť" in asyncio.run(h.t.undo_last())


def test_correction_down_via_table_unticks_reminder(h):
    h.tick_at(at(6, 0))
    h.text(at(7, 5), "6"); h.tick_at(at(7, 6))
    assert h.reminder("Ráno").completed
    h.edit_table(D, Ráno=2)
    h.tick_at(at(7, 9))
    assert not h.reminder("Ráno").completed and h.reminder("Ráno").summary == "💪 Ráno: 4 klikov"


def test_extra_reps_after_full_morning_go_to_evening(h):
    h.tick_at(at(6, 0))
    r1 = h.text(at(7, 5), "5")
    assert "Ráno: 5/6" in r1
    r2 = h.text(at(7, 30), "1")            # ráno 6/6 → hotové
    assert "Ráno: 6/6" in r2
    r3 = h.text(at(13, 45), "5")           # poobede, ranné vedro plné → večer
    assert "Večer: 5/6" in r3 and "Dnes: 11/12" in r3
    day = h.store.get_day(D)
    assert (day.morning, day.evening) == (6, 5)
    # explicitné „ráno“ stále funguje
    assert "Ráno: 8/6" in h.text(at(14, 0), "2 ráno")


# ── iCloud Pripomienky cez iOS Skratku ──────────────────────────────────────

def test_shortcut_plan_and_tick_off(hs):
    hs.tick_at(at(6, 0))
    plan = hs.t.shortcut_plan()
    assert plan["zoznam"] == "Kliky" and plan["pocet"] == 2 and plan["ciel"] == 12
    m, e = plan["pripomienky"]
    assert m["nazov"] == "💪 Ráno: 6 klikov" and m["cas_kratky"] == "07:00"
    assert m["poznamka"] == "kliky:2026-09-02:morning"
    assert e["nazov"] == "💪 Večer: 6 klikov" and e["cas"].endswith("19:20:00+02:00")

    # telefón hlási odškrtnutie rannej
    assert hs.t.shortcut_report({"kind": "done", "poznamka": m["poznamka"]}) == {"ok": True}
    hs.tick_at(at(8, 0))
    assert hs.store.get_day(D).morning == 6
    msgs = hs.take()
    assert len(msgs) == 1 and msgs[0].startswith("📱")
    # v pláne už ranná nie je, večerná ukazuje zvyšok
    plan = hs.t.shortcut_plan()
    assert plan["pocet"] == 1 and plan["pripomienky"][0]["nazov"] == "💪 Večer: 6 klikov"
    # opakované hlásenie to nepripíše druhýkrát
    hs.t.shortcut_report({"kind": "done", "poznamka": m["poznamka"]})
    hs.tick_at(at(8, 5))
    assert hs.store.get_day(D).morning == 6
    assert hs.table_rows()[D][1] == 6      # zapísané aj do tabuľky


def test_shortcut_completes_whole_day(hs):
    hs.tick_at(at(6, 0))
    for r in hs.t.shortcut_plan()["pripomienky"]:
        hs.t.shortcut_report({"kind": "done", "poznamka": r["poznamka"]})
    hs.tick_at(at(20, 0))
    day = hs.store.get_day(D)
    assert (day.morning, day.evening, day.total) == (6, 6, 12) and day.done
    assert any("🎉" in m for m in hs.take())
    assert hs.t.shortcut_plan()["pocet"] == 0


def test_shortcut_learns_edited_title_and_time(hs):
    hs.tick_at(at(6, 0))
    note = "kliky:2026-09-02:morning"
    hs.t.shortcut_report({"kind": "edit", "poznamka": note, "nazov": "Kliky ráno – 6 kusov",
                          "cas": "06:30"})
    hs.tick_at(at(5, 0))
    s = hs.t.settings()
    assert s.morning_title == "Kliky ráno – {n} kusov" and s.morning_time == "06:30"
    assert any("podľa tvojej úpravy" in m for m in hs.take())
    plan = hs.t.shortcut_plan()
    assert plan["pripomienky"][0]["nazov"] == "Kliky ráno – 6 kusov"
    assert plan["pripomienky"][0]["cas_kratky"] == "06:30"
    # nový čas platí aj pre výzvy v chate
    hs.tick_at(at(6, 30))
    assert [m.startswith("☀️") for m in hs.take()] == [True]


def test_shortcut_deleted_reminder_not_offered_again_today(hs):
    hs.tick_at(at(6, 0))
    hs.t.shortcut_report({"kind": "deleted", "poznamka": "kliky:2026-09-02:evening"})
    hs.tick_at(at(6, 5))
    plan = hs.t.shortcut_plan()
    assert [r["faza"] for r in plan["pripomienky"]] == ["morning"]
    # zajtra normálne – vo štvrtok je v pláne ráno (večer tam nie je vôbec)…
    hs.tick_at(at(0, 0, 30, D + timedelta(days=1)))
    assert [r["faza"] for r in hs.t.shortcut_plan()["pripomienky"]] == ["morning"]
    # …a v piatok zas obe, takže stredajšie zmazanie naozaj platilo len na stredu
    hs.tick_at(at(0, 0, 30, D + timedelta(days=2)))
    assert [r["faza"] for r in hs.t.shortcut_plan()["pripomienky"]] == ["morning", "evening"]


def test_shortcut_plan_empty_when_frozen(hs):
    hs.tick_at(at(6, 0))
    asyncio.run(hs.t.set_frozen(True))
    hs.tick_at(at(6, 5))
    plan = hs.t.shortcut_plan()
    assert plan["zamrazene"] is True and plan["pocet"] == 0
    asyncio.run(hs.t.set_frozen(False))
    hs.tick_at(at(6, 10))
    assert hs.t.shortcut_plan()["pocet"] == 2


def test_shortcut_ignores_foreign_and_old_notes(hs):
    hs.tick_at(at(6, 0))
    for bad in ({"kind": "done", "poznamka": "nakup mlieko"},
                {"kind": "done", "poznamka": "kliky:2026-08-31:morning"},
                {"kind": "cosi", "poznamka": "kliky:2026-09-02:morning"}):
        hs.t.shortcut_report(bad)
    hs.tick_at(at(6, 5))
    assert hs.store.get_day(D).total == 0 and hs.take() == []


# ── API pre mobilnú appku „Kliky" ───────────────────────────────────────────

def test_app_api_status_and_reporting(h):
    h.tick_at(at(6, 0))
    st = h.t.shortcut_plan()["stav"]
    assert st["ciel"] == 12 and st["rano_ciel"] == 6 and st["vecer_ciel"] == 6
    assert st["rano_cas"] == "07:00" and st["vecer_cas"] == "19:20"
    assert st["rano_due"].endswith("07:00:00+02:00") and st["vecer_due"].endswith("19:20:00+02:00")
    assert st["rano_nazov"] == "💪 Ráno: 6 klikov" and st["poznamka_rano"] == "kliky:2026-09-02:morning"
    assert st["zamrazene"] is False and st["splneny"] is False and st["zostava"] == 12

    # appka hlási kliky
    assert h.t.shortcut_report({"kind": "reps", "poznamka": st["poznamka_rano"], "n": "4"}) == {"ok": True}
    h.tick_at(at(7, 30))
    assert h.store.get_day(D).morning == 4
    assert h.t.shortcut_plan()["stav"]["zostava"] == 8
    assert h.table_rows()[D][1] == 4                     # zapísané aj do tabuľky

    # oprava na presnú hodnotu
    h.t.shortcut_report({"kind": "reps", "poznamka": st["poznamka_rano"], "n": "6", "absolute": True})
    h.tick_at(at(7, 35))
    st2 = h.t.shortcut_plan()["stav"]
    assert h.store.get_day(D).morning == 6 and st2["rano_hotovo"] is True
    assert st2["rano_nazov"] == "💪 Ráno: 6 klikov"      # splnené: koľko si dal, nie 0


def test_app_api_freeze_roundtrip(h):
    h.tick_at(at(6, 0))
    h.t.shortcut_report({"kind": "freeze", "hodnota": "1"})
    h.tick_at(at(6, 2))
    assert h.t.settings().frozen and h.t.shortcut_plan()["stav"]["zamrazene"] is True
    assert any("Zamrazené" in m for m in h.take())
    h.tick_at(at(7, 0))
    assert h.take() == []                                 # zamrazené: žiadne výzvy
    h.t.shortcut_report({"kind": "freeze", "hodnota": "0"})
    h.tick_at(at(7, 2))
    assert not h.t.settings().frozen and any("Odmrazené" in m for m in h.take())


def test_app_api_completing_the_day(h):
    h.tick_at(at(6, 0))
    st = h.t.shortcut_plan()["stav"]
    h.t.shortcut_report({"kind": "reps", "poznamka": st["poznamka_rano"], "n": "6"})
    h.t.shortcut_report({"kind": "reps", "poznamka": st["poznamka_vecer"], "n": "6"})
    h.tick_at(at(20, 0))
    day = h.store.get_day(D)
    assert (day.morning, day.evening) == (6, 6) and day.done
    assert any("🎉" in m for m in h.take())
    st = h.t.shortcut_plan()["stav"]
    # streak sa po novom ráta po týždňoch – jedna splnená streda ho ešte nedvíha
    assert st["splneny"] and st["streak"] == 0 and st["zajtra_ciel"] == 6


# ── týždenný plán: prechody dní, X a ticho vo voľných fázach ─────────────────

def test_nedela_prechod_na_pondelok_dvihne_x(tmp_path):
    """Nedeľa je voľno: žiadny ❌ verdikt a v pondelok nový riadok 4X s vyšším X."""
    h = Harness(tmp_path, at(6, 0, 0, D_NE))
    h.tick_at(at(6, 0, 0, D_NE))
    nedela = h.store.get_day(D_NE)
    assert nedela.goal == 0 and nedela.is_rest and h.t.settings().x_for(D_NE) == 3
    h.tick_at(at(22, 0, 0, D_NE))
    h.take()
    h.tick_at(at(0, 0, 30, D_PO))
    po = h.store.get_day(D_PO)
    assert h.t.settings().x_for(D_PO) == 4                    # každý pondelok +x_step
    assert po.goal == 16 and (po.morning_target, po.evening_target) == (8, 8)
    assert h.take() == []                                     # za nedeľu sa nenadáva
    assert h.table_rows()[D_PO][0] == 16
    assert h.table_rows()[D_NE][3].startswith("🌙")           # nedeľa ostáva „voľno“


def test_x_nenarastie_dvakrat_ked_bot_v_pondelok_restartuje(tmp_path):
    """X je odvodené z dátumu, nie počítadlo – tri reštarty v pondelok ho nezdvihnú."""
    h = Harness(tmp_path, at(6, 0, 0, D_NE))
    h.tick_at(at(20, 0, 0, D_NE))
    h.tick_at(at(0, 0, 30, D_PO))
    assert h.store.get_day(D_PO).goal == 16
    for kedy in (at(0, 5, 0, D_PO), at(6, 0, 0, D_PO), at(6, 40, 0, D_PO)):
        h.restart(kedy)
        h.tick_at(kedy)
    s = h.t.settings()
    assert (s.x, s.x_since, s.x_step) == (3, "2026-08-31", 1)
    assert s.x_for(D_PO) == 4 and h.store.get_day(D_PO).goal == 16
    assert h.table_rows()[D_PO][0] == 16


def test_ciel_uprostred_tyzdna_prepocita_dnesok_aj_zvysok(h):
    """/ciel 5 v stredu: dnešok hneď na 4X, budúce riadky tiež a zvyšok týždňa z plánu."""
    h.tick_at(at(6, 0))
    h.store.save_day(Day(D + timedelta(days=2), 12))          # piatkový riadok už v DB
    msg = asyncio.run(h.t.set_x(5))
    assert "X = 5 od tohto týždňa" in msg and "Dnes (streda): 10 ráno + 10 večer" in msg
    assert "štvrtok: 10 ráno" in msg and "nedeľa: voľno 🌙" in msg
    assert h.store.get_day(D).goal == 20                       # dnešok sa prepočítal hneď
    assert h.store.get_day(D + timedelta(days=2)).goal == 20    # aj budúci riadok
    h.tick_at(at(6, 5))
    assert h.table_rows()[D][0] == 20
    h.tick_at(at(0, 0, 30, D + timedelta(days=1)))
    assert h.store.get_day(D + timedelta(days=1)).goal == 10    # štvrtok: len ráno 2X
    h.take()
    h.tick_at(at(0, 0, 30, D_PO))
    assert h.store.get_day(D_PO).goal == 24                    # nový týždeň: X = 6


def test_ciel_v_nedelu_plati_az_od_nasledujuceho_pondelka(tmp_path):
    h = Harness(tmp_path, at(6, 0, 0, D_NE))
    h.tick_at(at(6, 0, 0, D_NE))
    msg = asyncio.run(h.t.set_x(9))
    assert "od pondelka 7.9." in msg and "pondelok: 18 ráno + 18 večer" in msg
    assert "Dnes (" not in msg                                 # v nedeľu niet čo meniť
    assert h.t.settings().x_since == "2026-09-07"
    assert h.store.get_day(D_NE).goal == 0                     # nedeľa ostáva voľno
    h.tick_at(at(0, 0, 30, D_PO))
    assert h.store.get_day(D_PO).goal == 36 and h.t.settings().x_for(D_PO) == 9


def test_v_nedelu_nepride_ziadna_vyzva_ani_pripomienka(tmp_path):
    h = Harness(tmp_path, at(6, 0, 0, D_NE))
    for hh in range(6, 24):
        h.tick_at(at(hh, 0, 0, D_NE))
        h.tick_at(at(hh, 40, 0, D_NE))
    assert h.sent == []                                        # celý deň ticho
    assert h.todos.list() == []                                # ani pripomienky v iCloude
    assert h.table_rows()[D_NE][0] == 0 and h.table_rows()[D_NE][3].startswith("🌙")
    # kliky v nedeľu sa zapíšu ako bonus, deň ostáva voľnom
    reply = h.text(at(10, 0, 0, D_NE), "5")
    assert "voľno" in reply and "bonus" in reply
    den = h.store.get_day(D_NE)
    assert den.total == 5 and den.is_rest and not den.done


def test_utorok_vecer_ziadna_vyzva(tmp_path):
    h = Harness(tmp_path, at(6, 0, 0, D_UT))
    h.tick_at(at(6, 0, 0, D_UT))
    assert h.store.get_day(D_UT).goal == 6                     # utorok = len ráno 2X
    h.tick_at(at(7, 0, 0, D_UT))
    assert [m.startswith("☀️") for m in h.take()] == [True]    # ráno výzva áno
    for kedy in (at(19, 20, 0, D_UT), at(20, 0, 0, D_UT), at(21, 30, 0, D_UT)):
        h.tick_at(kedy)
    assert h.take() == []                                      # večer už nič
    assert [i.summary for i in h.todos.list()] == ["💪 Ráno: 6 klikov"]


def test_utorok_kliky_o_20_sa_nezapisu_do_rana(tmp_path):
    """V utorok večerná fáza nie je – kliky o 20:00 patria večeru (bonus), nie ránu."""
    h = Harness(tmp_path, at(6, 0, 0, D_UT))
    h.tick_at(at(6, 0, 0, D_UT))
    reply = h.text(at(20, 0, 0, D_UT), "3")
    den = h.store.get_day(D_UT)
    assert (den.morning, den.evening) == (0, 3)
    assert "navyše" in reply and "Dnes: 3/6" in reply


def test_vypadok_bota_od_piatka_do_pondelka(tmp_path):
    """Zameškané tréningové dni sú zamrazené, nedeľa ostáva voľnom a v pondelok
    beží nový týždeň – bez dobiehania starých výziev a bez ❌ za voľnú nedeľu."""
    d_pi, d_so = date(2026, 9, 4), date(2026, 9, 5)
    h = Harness(tmp_path, at(6, 0, 0, d_pi))
    h.tick_at(at(6, 0, 0, d_pi))
    h.text(at(8, 0, 0, d_pi), "12")                            # piatok splnený
    h.take()
    h.tick_at(at(11, 0, 0, D_PO))                              # prvý tick až v pondelok
    days = {x.date: x for x in h.store.all_days()}
    assert days[d_so].frozen is True                           # sobota: bot nebežal, nie je to chyba
    assert days[D_NE].frozen is False and days[D_NE].is_rest    # nedeľu netreba mraziť
    assert days[D_PO].goal == 16 and days[D_PO].frozen is False
    msgs = h.take()
    assert not any(m.startswith("❌") for m in msgs)
    assert not any(m.startswith("☀️") for m in msgs)           # výzva spred 4 h sa nedobieha
    assert sorted(h.table_rows()) == [d_pi, d_so, D_NE, D_PO]


def test_cely_splneny_tyzden_zdvihne_streak(tmp_path):
    """Streak = celé splnené týždne: pondelok–sobota splnené → v nedeľu streak 1."""
    d_po = date(2026, 8, 31)
    h = Harness(tmp_path, at(6, 0, 0, d_po))
    h.tick_at(at(6, 0, 0, d_po))
    for i in range(6):                                         # pondelok až sobota
        d = d_po + timedelta(days=i)
        h.text(at(8, 0, 0, d), str(h.store.get_day(d).goal))
        h.tick_at(at(0, 0, 30, d + timedelta(days=1)))
        assert not any(m.startswith("❌") for m in h.take())
    h.tick_at(at(9, 0, 0, D_NE))
    st = h.t.shortcut_plan()["stav"]
    assert st["volno"] is True and st["streak"] == 1
    h.tick_at(at(0, 0, 30, D_PO))
    assert h.t.shortcut_plan()["stav"]["streak"] == 1          # nový týždeň streak nezhodí


def test_app_status_v_utorok_a_v_nedelu(tmp_path):
    hu = Harness(tmp_path / "ut", at(6, 0, 0, D_UT))
    hu.tick_at(at(6, 0, 0, D_UT))
    st = hu.t.shortcut_plan()["stav"]
    assert st["den"] == "utorok" and st["x"] == 3 and st["volno"] is False
    assert (st["ciel"], st["rano_ciel"], st["vecer_ciel"]) == (6, 6, 0)
    assert st["rano_due"].endswith("07:00:00+02:00") and st["vecer_due"] is None
    assert st["poznamka_rano"] == "kliky:2026-09-01:morning" and st["poznamka_vecer"] == ""
    assert st["vecer_hotovo"] is True and st["stav"] == "open"
    assert st["zajtra_ciel"] == 12                             # streda: 6 ráno + 6 večer

    hn = Harness(tmp_path / "ne", at(6, 0, 0, D_NE))
    hn.tick_at(at(6, 0, 0, D_NE))
    st = hn.t.shortcut_plan()["stav"]
    assert st["den"] == "nedeľa" and st["volno"] is True and st["stav"] == "rest"
    assert (st["ciel"], st["zostava"]) == (0, 0)
    assert st["rano_due"] is None and st["vecer_due"] is None
    assert st["poznamka_rano"] == "" and st["poznamka_vecer"] == ""
    assert st["rano_hotovo"] is True and st["vecer_hotovo"] is True
    assert st["zajtra_ciel"] == 16                             # pondelok už s novým X


def test_bol_to_ciel_po_splnenom_dni_neumlci_dalsie_splnenie(h):
    """Po „🎯 Bol to X“ deň splnený nie je – keď ho naozaj splníš, musí prísť 🎉."""
    h.tick_at(at(6, 0))
    assert "🎉" in h.text(at(7, 5), "12")
    h.take()
    msg = asyncio.run(h.t.undo_last(as_goal=True))
    assert "X = 12" in msg
    den = h.store.get_day(D)
    assert den.goal == 48 and not den.done
    h.edit_table(D, Ráno=24, Večer=24)
    h.tick_at(at(7, 20))
    assert h.store.get_day(D).done
    assert any("🎉" in m for m in h.take())
