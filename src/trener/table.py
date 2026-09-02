"""Tabuľka „kliky.xlsx“ – zdroj pravdy pre dni a nastavenia.

Hárok „Kliky“: Dátum | Cieľ | Ráno | Večer | Spolu | Stav | Streak | Poznámka
  - stĺpce sa hľadajú PODĽA HLAVIČKY (používateľ si môže stĺpce presúvať a pridávať vlastné),
  - Cieľ, Ráno, Večer, Poznámka: píše používateľ aj bot (3-cestný merge, tabuľka vyhráva),
  - Spolu, Stav, Streak: počíta bot; do „Stav“ môže používateľ napísať „zamrazený“,
  - bot upravuje hárok NA MIESTE: cudzie riadky, stĺpce, komentáre a vzorce nechá tak,
    nové dni dopíše na koniec; riadok, ktorý používateľ zmazal, zmaže aj u seba
    (okrem dnešného a prípadu, keď je tabuľka náhle celá prázdna).
Hárok „Nastavenia“: názov | hodnota | popis (rovnaký merge). Hárok „Návod“: text.

Všetko tu pracuje s bajtmi (BytesIO) – žiadny prístup na sieť, plne testovateľné.
"""
from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time as dtime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from trener.model import (DONE, FAILED, FROZEN, FROZEN_WORDS, OPEN, STATUS_LABEL, Day, Settings,
                          parse_hhmm)

SHEET_DAYS = "Kliky"
SHEET_SETTINGS = "Nastavenia"
SHEET_HELP = "Návod"
DAY_HEADERS = ["Dátum", "Cieľ", "Ráno", "Večer", "Spolu", "Stav", "Streak", "Poznámka"]
KEYS = ["date", "goal", "morning", "evening", "total", "status", "streak", "note"]
HEADER_KEY = dict(zip(DAY_HEADERS, KEYS))
DEFAULT_LAYOUT = {k: i + 1 for i, k in enumerate(KEYS)}      # 1-based
COL = {h: i + 1 for i, h in enumerate(DAY_HEADERS)}          # spätná kompatibilita (testy)
USER_COLS = ("goal", "morning", "evening")

# (kľúč v Settings, popisok v tabuľke, popis)
SETTINGS_ROWS = [
    ("frozen", "Zamrazené", "ÁNO = tréner nič nepripomína (chat ani Pripomienky), tabuľku ďalej sleduje."),
    ("increment", "Prírastok cieľa", "O koľko klikov rastie cieľ po každom splnenom dni (0 = nerastie)."),
    ("morning_time", "Ranný čas", "Kedy príde ranná pripomienka a prvá výzva v chate (HH:MM)."),
    ("evening_time", "Večerný čas", "Kedy príde večerná pripomienka a prvá výzva v chate (HH:MM)."),
    ("nag_max", "Max výziev v chate", "Koľko výziev za sebou pošle bot v jednej fáze (ráno / večer)."),
    ("nag_interval_min", "Rozostup výziev (min)", "Minúty medzi výzvami v chate."),
    ("morning_title", "Názov rannej pripomienky", "Šablóna; {n} = počet klikov. Zmena názvu v appke Pripomienky sa sem prepíše sama."),
    ("evening_title", "Názov večernej pripomienky", "Šablóna; {n} = počet klikov."),
    ("summary_on_fail", "Správa pri nesplnenom dni", "ÁNO = o polnoci príde jedna správa, ak deň nebol splnený."),
]
LABEL_TO_KEY = {label: key for key, label, _ in SETTINGS_ROWS}
BOOL_KEYS = {"frozen", "summary_on_fail"}
INT_KEYS = {"increment", "nag_max", "nag_interval_min"}
TIME_KEYS = {"morning_time", "evening_time"}

HELP_TEXT = [
    "Virtuálny tréner klikov – ako funguje táto tabuľka",
    "",
    "• Hárok „Kliky“: jeden riadok = jeden deň. Do stĺpcov Cieľ, Ráno, Večer a Poznámka môžeš písať ty aj bot.",
    "• Ráno = počet klikov, ktoré si dal v rannej fáze, Večer = vo večernej. Spolu, Stav a Streak počíta bot.",
    "• Deň je splnený, keď Ráno + Večer ≥ Cieľ. Rozdelenie cieľa je polovica ráno (zaokrúhlená nahor) a zvyšok večer.",
    "• Ak do stĺpca Stav napíšeš „zamrazený“, deň sa nepočíta a neprerušuje streak (napr. choroba).",
    "• Bot číta tabuľku každé 2 minúty. Čo zapíšeš sem, platí – nemusíš mu nič písať do chatu.",
    "• Stĺpce môžeš presúvať a pridávať vlastné (bot ich hľadá podľa názvu v prvom riadku a cudzie nechá tak).",
    "• Riadok, ktorý zmažeš, zmaže bot aj u seba (dnešný riadok si vždy dorobí).",
    "• Keď je v hárku Nastavenia „Zamrazené“ = ÁNO, bot nič nepripomína (chat ani Pripomienky), ale tabuľku ďalej sleduje.",
    "• Časy zadávaj ako HH:MM (napr. 07:00). Zmena času pripomienky v appke Pripomienky sa sem prepíše sama.",
    "• Súbor ukladaj v Exceli priamo sem (na sieťový disk). Kým ho máš otvorený, bot počká a zapíše, keď ho zavrieš.",
    "• Numbers na iPhone/Macu súbor otvorí, ale ukladá kópiu – zmeny radšej píš cez Excel, chat alebo Pripomienky.",
]


class TableError(Exception):
    pass


class TableCorrupt(TableError):
    """Súbor sa nedá prečítať (rozpísaný / poškodený) – NIKDY ho neprepisuj."""


@dataclass
class RowVals:
    date: date
    goal: int | None
    morning: int | None
    evening: int | None
    note: str
    status_text: str
    frozen_by_user: bool
    formula_cells: set[str] = field(default_factory=set)   # 'goal'/'morning'/'evening' so vzorcom
    extra: dict[int, object] = field(default_factory=dict) # cudzie stĺpce (index → hodnota)
    row: int = 0                                           # riadok v hárku
    total_text: str = ""


@dataclass
class Parsed:
    days: dict[date, RowVals]
    settings: dict[str, object]     # kľúč Settings → hodnota (už pretypovaná)
    warnings: list[str]
    has_days_sheet: bool
    has_settings_sheet: bool
    layout: dict[str, int] = field(default_factory=lambda: dict(DEFAULT_LAYOUT))
    max_row: int = 1
    foreign_rows: int = 0           # riadky, ktorým nerozumieme (nechávame ich tak)


# ── konverzie buniek ─────────────────────────────────────────────────────────

def _norm(s) -> str:
    s = "" if s is None else str(s)
    return "".join(c for c in unicodedata.normalize("NFD", s.strip().lower()) if unicodedata.category(c) != "Mn")


def cell_int(v) -> int | None:
    """Číslo z bunky: 5, 5.0, '5', ' 5 ', '5 klikov' → 5; prázdne → None."""
    if v is None:
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(round(v))
    s = str(v).strip()
    if not s:
        return None
    m = re.search(r"-?\d+", s)
    return int(m.group()) if m else None


def cell_date(v) -> date | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    if not s:
        return None
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        y, mo, d = map(int, m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2})\.\s*(\d{1,2})\.\s*(\d{4})", s)
        if not m:
            return None
        d, mo, y = map(int, m.groups())
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def cell_bool(v) -> bool | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return bool(v)
    s = _norm(v)
    if s in ("ano", "yes", "y", "true", "1", "x", "✓", "✔", "da"):
        return True
    if s in ("nie", "no", "n", "false", "0", "", "-"):
        return False
    return None


def cell_time(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%H:%M")
    if isinstance(v, dtime):
        return v.strftime("%H:%M")
    if isinstance(v, float) and 0 <= v < 1:   # Excel zlomok dňa
        mins = int(round(v * 24 * 60))
        return f"{mins // 60:02d}:{mins % 60:02d}"
    return parse_hhmm(str(v))


def bool_label(b: bool) -> str:
    return "ÁNO" if b else "NIE"


def _is_frozen_text(status_text: str) -> bool:
    n = _norm(status_text)
    return n in {_norm(w) for w in FROZEN_WORDS} or any(w in n for w in ("zamraz", "frozen", "❄"))


# ── hlavička ─────────────────────────────────────────────────────────────────

def _detect_layout(ws) -> tuple[dict[str, int], bool]:
    """Mapa kľúč → index stĺpca podľa prvého riadku. Vráti (layout, hlavička_nájdená)."""
    norm_to_key = {_norm(h): k for h, k in HEADER_KEY.items()}
    layout: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        key = norm_to_key.get(_norm(ws.cell(1, c).value))
        if key and key not in layout:
            layout[key] = c
    found = "date" in layout
    if not found:
        return dict(DEFAULT_LAYOUT), False
    return layout, True


# ── čítanie ──────────────────────────────────────────────────────────────────

def parse_workbook(data: bytes) -> Parsed:
    try:
        wb_v = load_workbook(io.BytesIO(data), data_only=True)
        wb_f = load_workbook(io.BytesIO(data), data_only=False)
    except Exception as e:  # noqa: BLE001 – openpyxl vie hodiť všeličo pri rozpísanom súbore
        raise TableCorrupt(f"xlsx sa nedá načítať: {e}") from e
    warnings: list[str] = []
    days: dict[date, RowVals] = {}
    layout = dict(DEFAULT_LAYOUT)
    max_row = 1
    foreign = 0
    has_days = SHEET_DAYS in wb_v.sheetnames
    if has_days:
        ws_v, ws_f = wb_v[SHEET_DAYS], wb_f[SHEET_DAYS]
        layout, _ = _detect_layout(ws_f)
        max_row = ws_v.max_row
        known_cols = set(layout.values())
        for r in range(2, ws_v.max_row + 1):
            cells = [ws_v.cell(r, c).value for c in range(1, ws_v.max_column + 1)]
            if all(v is None or str(v).strip() == "" for v in cells):
                continue                                   # prázdny riadok
            raw_date = ws_v.cell(r, layout["date"]).value
            d = cell_date(raw_date)
            if d is None:
                warnings.append(f"riadok {r}: nerozumiem dátumu {raw_date!r} – nechávam ho tak")
                foreign += 1
                continue
            if d in days:
                warnings.append(f"riadok {r}: dátum {d} je v tabuľke dvakrát – beriem prvý (riadok {days[d].row})")
                foreign += 1
                continue
            formulas = set()
            vals = {}
            for name in USER_COLS:
                c = layout.get(name)
                if c is None:
                    vals[name] = None
                    continue
                fv = ws_f.cell(r, c).value
                if isinstance(fv, str) and fv.startswith("="):
                    formulas.add(name)
                vals[name] = cell_int(ws_v.cell(r, c).value)
            sv = ws_v.cell(r, layout["status"]).value if "status" in layout else None
            status_text = "" if sv is None else str(sv).strip()
            tv = ws_v.cell(r, layout["total"]).value if "total" in layout else None
            total_text = "" if tv is None else str(tv).strip()          # 0 je platná hodnota, nie „prázdne“
            note = ws_v.cell(r, layout["note"]).value if "note" in layout else None
            extra = {c: ws_f.cell(r, c).value for c in range(1, ws_v.max_column + 1)
                     if c not in known_cols and ws_f.cell(r, c).value is not None}
            days[d] = RowVals(d, vals["goal"], vals["morning"], vals["evening"],
                              "" if note is None else str(note), status_text, _is_frozen_text(status_text),
                              formulas, extra, r, total_text)
    settings: dict[str, object] = {}
    has_settings = SHEET_SETTINGS in wb_v.sheetnames
    if has_settings:
        ws = wb_v[SHEET_SETTINGS]
        for r in range(1, ws.max_row + 1):
            label = ws.cell(r, 1).value
            if label is None:
                continue
            key = LABEL_TO_KEY.get(str(label).strip())
            if key is None:
                continue
            v = ws.cell(r, 2).value
            parsed = _parse_setting(key, v)
            if parsed is None:
                if v not in (None, ""):
                    warnings.append(f"Nastavenia „{label}“: nerozumiem hodnote {v!r}")
                continue
            settings[key] = parsed
    return Parsed(days, settings, warnings, has_days, has_settings, layout, max_row, foreign)


def _parse_setting(key: str, v):
    if key in BOOL_KEYS:
        return cell_bool(v)
    if key in INT_KEYS:
        n = cell_int(v)
        return n if n is not None and n >= 0 else None
    if key in TIME_KEYS:
        return cell_time(v)
    if v is None:
        return None
    s = str(v).strip()
    return s or None


# ── merge ────────────────────────────────────────────────────────────────────

@dataclass
class MergeResult:
    days: list[Day]                  # výsledné dni (všetky)
    settings: Settings
    from_table: list[str]            # čo sme prevzali z tabuľky (na log / oznámenie)
    to_table: bool                   # treba tabuľku zapísať
    changed_days: list[date]
    changed_settings: list[str]
    deleted_days: list[date] = field(default_factory=list)   # používateľ zmazal riadok → zmazať aj u bota


def merge(parsed: Parsed | None, bot_days: list[tuple[Day, Day | None]], bot_settings: Settings,
          settings_synced: dict[str, str | None], today: date, default_goal: int = 10) -> MergeResult:
    """3-cestný merge: (tabuľka, bot, posledná zosynchronizovaná snímka).

    Pre každé pole: zmena v tabuľke voči snímke → vyhráva tabuľka; inak zmena bota → do tabuľky.
    Riadky navyše v tabuľke sa importujú; riadky, ktoré bot pozná ako zosynchronizované a v
    tabuľke už nie sú, používateľ zmazal (okrem dnešného a okrem prázdnej tabuľky).
    """
    from_table: list[str] = []
    changed_days: list[date] = []
    deleted: list[date] = []
    to_table = False
    result: dict[date, Day] = {}
    trows = parsed.days if parsed else {}
    bot_by_date = {d.date: (d, s) for d, s in bot_days}
    table_has_rows = bool(trows)
    all_dates = sorted(set(trows) | set(bot_by_date))

    def prev_goal(d: date) -> int:
        for x in reversed(all_dates):
            if x < d:
                src = result.get(x) or (bot_by_date.get(x) or (None,))[0]
                if src is not None and src.goal > 0:
                    return src.goal
                tv = trows.get(x)
                if tv is not None and tv.goal:
                    return tv.goal
        return default_goal

    for d in all_dates:
        t = trows.get(d)
        bot, synced = bot_by_date.get(d, (None, None))
        if t is None and bot is not None:
            if synced is not None and d != today and table_has_rows and parsed is not None and parsed.has_days_sheet:
                deleted.append(d)              # bol v tabuľke a už nie je → používateľ ho zmazal
                from_table.append(f"{d.day}.{d.month}.: riadok zmazaný (tabuľka)")
                continue
            result[d] = bot
            to_table = True                    # v tabuľke chýba → dopíšeme
            continue
        if bot is None and t is not None:
            goal = t.goal if t.goal else prev_goal(d)     # bez cieľa → posledný známy cieľ, nie 0
            day = Day(d, goal, t.morning or 0, t.evening or 0, t.frozen_by_user, t.note)
            result[d] = day
            from_table.append(f"{d.day}.{d.month}.: nový riadok {day.morning}+{day.evening}/{day.goal}")
            changed_days.append(d)
            to_table = True                    # dopočítané stĺpce (Spolu/Stav/Streak) treba zapísať
            continue
        assert t is not None and bot is not None
        new = bot
        row_to_table = False
        for name in USER_COLS:
            tv = getattr(t, name)
            bv = getattr(bot, name)
            sv = getattr(synced, name) if synced else None
            tv_eff = tv if tv is not None else 0
            if name == "goal" and tv is None:
                row_to_table = True          # prázdny cieľ = nezadaný → ostáva hodnota bota, dopíšeme ju
                continue
            if name in t.formula_cells:
                # bunka so vzorcom patrí používateľovi – berieme jej hodnotu, nikdy ju neprepisujeme
                if tv is not None and tv != bv:
                    new = new.copy(**{name: tv})
                    from_table.append(f"{d.day}.{d.month}.: {_sk(name)} = {tv} (vzorec)")
                continue
            if synced is None or tv_eff != (sv if sv is not None else 0):
                if tv_eff != bv:
                    new = new.copy(**{name: tv_eff})
                    from_table.append(f"{d.day}.{d.month}.: {_sk(name)} {bv} → {tv_eff} (tabuľka)")
                if tv is None:
                    row_to_table = True     # prázdna bunka → dopíšeme 0
            elif bv != sv:
                row_to_table = True         # zmena bota → do tabuľky
        # poznámka
        tn, bn, sn = t.note or "", bot.note or "", (synced.note if synced else None)
        if synced is None or tn != (sn or ""):
            if tn != bn:
                new = new.copy(note=tn)
        elif bn != sn:
            row_to_table = True
        # zamrazený deň zo stĺpca Stav
        if t.frozen_by_user and not bot.frozen:
            if synced is not None and synced.frozen:
                row_to_table = True          # náš starý zápis, bot medzitým odmrazil → vyhráva bot
            else:
                new = new.copy(frozen=True)
                from_table.append(f"{d.day}.{d.month}.: zamrazený (tabuľka)")
        elif not t.frozen_by_user and bot.frozen and synced is not None and synced.frozen \
                and t.status_text and not new.done:
            new = new.copy(frozen=False)
            from_table.append(f"{d.day}.{d.month}.: odmrazený (tabuľka)")
            row_to_table = True
        elif bot.frozen != (synced.frozen if synced else None):
            row_to_table = True
        if new != bot:
            changed_days.append(d)
            row_to_table = True              # používateľ menil hodnoty → prepočítať Spolu/Stav/Streak
        result[d] = new
        if row_to_table or _computed_stale(t, new, today):
            to_table = True

    # nastavenia
    settings = bot_settings.copy()
    changed_settings: list[str] = []
    tset = parsed.settings if parsed else {}
    for key, _label, _ in SETTINGS_ROWS:
        bv = getattr(bot_settings, key)
        sv_raw = settings_synced.get(key)
        sv = _from_snapshot(key, sv_raw) if sv_raw is not None else None
        if key in tset:
            tv = tset[key]
            if sv is None or tv != sv:
                if tv != bv:
                    setattr(settings, key, tv)
                    changed_settings.append(key)
                    from_table.append(f"nastavenie {_label}: {bv} → {tv} (tabuľka)")
            elif bv != sv:
                to_table = True
        else:
            to_table = True   # v tabuľke chýba → dopíšeme
    if parsed is None or not parsed.has_days_sheet or not parsed.has_settings_sheet:
        to_table = True
    return MergeResult(sorted(result.values(), key=lambda x: x.date), settings, from_table,
                       to_table, changed_days, changed_settings, deleted)


def _sk(name: str) -> str:
    return {"goal": "cieľ", "morning": "ráno", "evening": "večer"}[name]


def _from_snapshot(key: str, s: str):
    if key in BOOL_KEYS:
        return s in ("1", "true", "True")
    if key in INT_KEYS:
        try:
            return int(s)
        except ValueError:
            return None
    return s


def _computed_stale(t: RowVals, day: Day, today: date) -> bool:
    if t.status_text.strip() != STATUS_LABEL[day.status(today)]:
        return True
    return cell_int(t.total_text) != day.total


# ── zápis ────────────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
BOT_FILL = PatternFill("solid", fgColor="F2F2F2")
STATUS_FILL = {DONE: "E2F0D9", FAILED: "FBE5D6", FROZEN: "DEEBF7", OPEN: "FFF2CC"}


def _set_text(cell, value) -> None:
    """Text do bunky tak, aby ho Excel nebral ako vzorec (poznámka „=fajn“)."""
    if value is None or value == "":
        cell.value = None
        return
    cell.value = str(value)
    if str(value).startswith("="):
        cell.data_type = "s"


def render_workbook(existing: bytes | None, days: list[Day], settings: Settings,
                    streaks: dict[date, int], today: date, parsed: Parsed | None = None,
                    extras: dict[date, dict[int, object]] | None = None) -> bytes:
    """Zapíše dni + nastavenia. `existing` sa upraví NA MIESTE podľa `parsed` (rozloženie
    stĺpcov, riadky dní); cudzie riadky/stĺpce/komentáre/vzorce ostávajú."""
    wb = None
    if existing:
        try:
            wb = load_workbook(io.BytesIO(existing))
        except Exception:  # noqa: BLE001
            wb = None
    if wb is None:
        wb = Workbook()
        wb.remove(wb.active)
        parsed = None
    fresh = SHEET_DAYS not in wb.sheetnames
    ws = wb[SHEET_DAYS] if not fresh else wb.create_sheet(SHEET_DAYS, 0)
    layout = dict(parsed.layout) if (parsed is not None and parsed.has_days_sheet) else dict(DEFAULT_LAYOUT)
    # chýbajúce známe stĺpce doplň na koniec (napr. používateľ zmazal „Streak“)
    if not fresh:
        _, found = _detect_layout(ws)
        if not found:
            layout = dict(DEFAULT_LAYOUT)
        next_col = max([ws.max_column] + list(layout.values())) + 1
        for key in KEYS:
            if key not in layout:
                layout[key] = next_col
                next_col += 1
    key_to_header = {k: h for h, k in HEADER_KEY.items()}
    for key, c in layout.items():
        cell = ws.cell(1, c)
        if cell.value is None or _norm(cell.value) != _norm(key_to_header[key]):
            cell.value = key_to_header[key]
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    rows_by_date = {d: rv.row for d, rv in (parsed.days.items() if parsed else [])}
    formula_by_date = {d: rv.formula_cells for d, rv in (parsed.days.items() if parsed else [])}
    next_row = max([1, ws.max_row] + list(rows_by_date.values())) + 1
    if fresh or parsed is None:
        next_row = 2
    for day in sorted(days, key=lambda x: x.date):
        r = rows_by_date.get(day.date)
        if r is None:
            r = next_row
            next_row += 1
            ws.cell(r, layout["date"], day.date).number_format = "DD.MM.YYYY"
        formulas = formula_by_date.get(day.date, set())
        st = day.status(today)
        for key, value in (("goal", day.goal), ("morning", day.morning), ("evening", day.evening)):
            if key in formulas:
                continue                                   # vzorec používateľa nechávame
            cell = ws.cell(r, layout[key])
            if cell.value != value:
                cell.value = value
        ws.cell(r, layout["total"], day.total).fill = BOT_FILL
        c = ws.cell(r, layout["status"], STATUS_LABEL[st])
        c.fill = PatternFill("solid", fgColor=STATUS_FILL[st])
        ws.cell(r, layout["streak"], streaks.get(day.date, 0)).fill = BOT_FILL
        _set_text(ws.cell(r, layout["note"]), day.note or None)
        for col, v in (extras or {}).get(day.date, {}).items():
            ws.cell(r, col, v)
    if fresh:
        widths = {"date": 12, "goal": 7, "morning": 7, "evening": 7, "total": 7, "status": 14, "streak": 8, "note": 40}
        for key, w in widths.items():
            ws.column_dimensions[ws.cell(1, layout[key]).column_letter].width = w
        ws.freeze_panes = "A2"

    # nastavenia
    wss = wb[SHEET_SETTINGS] if SHEET_SETTINGS in wb.sheetnames else wb.create_sheet(SHEET_SETTINGS)
    for i, h in enumerate(("Nastavenie", "Hodnota", "Popis"), 1):
        c = wss.cell(1, i, h)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
    label_rows: dict[str, int] = {}
    for r in range(2, wss.max_row + 1):
        lab = wss.cell(r, 1).value
        if lab is not None:
            label_rows[str(lab).strip()] = r
    next_row = max([1] + list(label_rows.values())) + 1
    wss.data_validations.dataValidation = []           # žiadne hromadenie validácií
    dv_bool = DataValidation(type="list", formula1='"ÁNO,NIE"', allow_blank=True)
    wss.add_data_validation(dv_bool)
    for key, label, desc in SETTINGS_ROWS:
        r = label_rows.get(label)
        if r is None:
            r = next_row
            next_row += 1
        wss.cell(r, 1, label).font = Font(bold=True)
        v = getattr(settings, key)
        if key in BOOL_KEYS:
            cell = wss.cell(r, 2, bool_label(bool(v)))
            dv_bool.add(cell)
        elif key in TIME_KEYS:
            wss.cell(r, 2, str(v))
        else:
            cell = wss.cell(r, 2)
            if isinstance(v, str):
                _set_text(cell, v)
            else:
                cell.value = v
        wss.cell(r, 3, desc)
    wss.column_dimensions["A"].width = 30
    wss.column_dimensions["B"].width = 26
    wss.column_dimensions["C"].width = 90

    # návod
    wsh = wb[SHEET_HELP] if SHEET_HELP in wb.sheetnames else wb.create_sheet(SHEET_HELP)
    for i, line in enumerate(HELP_TEXT, 1):
        c = wsh.cell(i, 1, line)
        if i == 1:
            c.font = Font(bold=True, size=13)
    wsh.column_dimensions["A"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
